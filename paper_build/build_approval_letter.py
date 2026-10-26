"""Letter requesting CodiSmile Dental Clinic's approval to conduct the study.

    python paper_build/build_approval_letter.py

Writes ToothWise_Clinic_Approval_Letter.docx and .pdf in the repo root, on
the same UDM letterhead, margins, and font as the capstone paper.
"""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_paper import BODY_W, FONT, LETTERHEAD, ROOT, add_runs, anchor_behind, fmt, setup_section, to_pdf

NAME = "ToothWise_Clinic_Approval_Letter"
DATE = "October 26, 2026"
HEAD = "DRA. ANGELICA SHALA C. CODILLA"
PROPONENTS = ["Brayan Axel L. Tubillas", "Keith Ashly V. Sales", "Reizen S. Arreola",
              "Robin C. Ponce", "Rizelle Carla N. Magpantay"]
ADVISER = "PROF. FRANCIS ARLANDO L. ATIENZA"

BODY = [
    "Greetings of peace!",
    "We, the undersigned students of the College of Computing Studies, Universidad de Manila, "
    "are currently conducting our capstone project entitled **“Toothwise: A YOLO-Based Multi-Condition "
    "Dental Screening System with KNN-Based Recommendation for CodiSmile Dental Clinic.”** The system "
    "analyzes still intraoral photographs to detect tartar, plaque, cavity, gum inflammation, and healthy "
    "tooth using the YOLOv8 and YOLOv12 object detection models, computes a risk classification, and "
    "generates a personalized prevention plan, while all findings remain subject to the review and final "
    "judgment of the clinic's licensed dentist.",
    "In line with this, we respectfully request your approval to conduct our study at CodiSmile Dental "
    "Clinic, specifically to:",
]
ITEMS = [
    "use CodiSmile Dental Clinic as the deployment site and clinical reference of the Toothwise system;",
    "demonstrate and pilot the system with the clinic's dentist, dental assistants, and receptionist, using "
    "sample dental images only; and",
    "consult the clinic's dentist on the system's detection results, risk classification, and prevention "
    "plans so that they reflect sound clinical practice.",
]
CLOSING = [
    "The study does not involve the clinic's patients. The ISO/IEC 25010 software quality evaluation has "
    "already been answered by 50 respondents from outside the clinic, so no patient photograph, record, or "
    "personal information from CodiSmile Dental Clinic will be collected. Any information shared with us will "
    "be kept strictly confidential and used solely for academic purposes, in accordance with the Data Privacy "
    "Act of 2012 (Republic Act No. 10173), and all activities will be scheduled so as not to disrupt the "
    "clinic's daily operations.",
    "Your approval will be a great help to the completion of our study. Thank you very much for your time "
    "and kind consideration.",
]


def line(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, bold=False, size=None):
    par = fmt(doc.add_paragraph(), align, 0, before, after, 1.0)
    add_runs(par, text, size=size, bold=bold)
    return par


def signature(cell, name, role, first=False, gap=26):
    par = fmt(cell.paragraphs[0] if first else cell.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, 0, gap, 0, 1.0)
    add_runs(par, "_" * 32)
    add_runs(fmt(cell.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 0, 1.0), f"**{name}**")
    add_runs(fmt(cell.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 0, 1.0), role)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    setup_section(sec)
    anchor_behind(sec.header.paragraphs[0].add_run(), LETTERHEAD, 8.5, 11)

    line(doc, DATE, after=14)
    line(doc, f"**{HEAD}**")
    for t in ("Head Dentist", "CodiSmile Dental Clinic",
              "NW Phase LRT Central Terminal, Arroceros St.", "Brgy. 659-A, Ermita, Manila City"):
        line(doc, t)
    line(doc, "Dear Dra. Codilla:", before=14, after=8)

    for t in BODY:
        fmt(add_runs(doc.add_paragraph(), t), indent=0.5 if t != BODY[0] else 0, after=6, line=1.15)
    for i, t in enumerate(ITEMS, 1):
        par = fmt(doc.add_paragraph(), left=0.85, hanging=0.3, after=3, line=1.15)
        par.paragraph_format.tab_stops.add_tab_stop(Inches(0.85))
        add_runs(par, f"{i}.\t{t}")
    for i, t in enumerate(CLOSING):
        par = fmt(add_runs(doc.add_paragraph(), t), indent=0.5, before=3, after=6, line=1.15)
        if i == len(CLOSING) - 1:
            par.paragraph_format.page_break_before = True   # request on page 1; thanks and signatures on page 2

    line(doc, "Respectfully yours,", before=6, after=0)

    # Proponents in two columns, adviser below.
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    left, right = t.cell(0, 0), t.cell(0, 1)
    for c in (left, right):
        c.width = Inches(BODY_W / 2)
    for i, n in enumerate(PROPONENTS):
        signature(left if i < 3 else right, n, "Proponent", first=(i in (0, 3)))

    line(doc, "Noted by:", before=24)
    signature(doc.add_table(rows=1, cols=1).cell(0, 0), ADVISER, "Capstone Adviser, College of Computing Studies", True)

    # Approval block for the clinic.
    line(doc, "Approved by:", before=24)
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    a, b = t.cell(0, 0), t.cell(0, 1)
    a.width, b.width = Inches(3.9), Inches(BODY_W - 3.9)
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    t._tbl.tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), (3.9, BODY_W - 3.9)):
        gc.set(qn("w:w"), str(int(w * 1440)))
    signature(a, HEAD, "Head Dentist, CodiSmile Dental Clinic", True)
    par = fmt(b.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT, 0, 40, 0, 1.0)
    par.paragraph_format.tab_stops.add_tab_stop(Inches(2.5), WD_TAB_ALIGNMENT.LEFT)
    add_runs(par, "Date: ______________")
    add_runs(fmt(b.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 0, 1.0), "☐ Approved")
    add_runs(fmt(b.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, 0, 2, 0, 1.0), "☐ Not approved")

    out = ROOT / f"{NAME}.docx"
    doc.save(out)
    pdf = to_pdf(out)
    import pymupdf
    print(f"Wrote {out.name} and {pdf.name}: {len(pymupdf.open(pdf))} page(s)")


if __name__ == "__main__":
    main()
