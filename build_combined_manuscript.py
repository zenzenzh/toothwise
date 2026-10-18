"""Render revised Chapters IV-V and append them to the supplied Chapters I-III PDF."""
from pathlib import Path

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfbase.pdfmetrics import stringWidth
from xml.sax.saxutils import escape

ROOT = Path(__file__).parent
DOCX = ROOT / 'ToothWise_Chapters_4_and_5_Revised.docx'
SOURCE = Path(r'C:\Users\arreo\Downloads\toothwise-updated-real-chapter-1-3-midtermpdfversion.pdf')
CHAPTER_PDF = ROOT / 'chapters_4_5_revised.pdf'
OUTPUT = ROOT / 'ToothWise_Complete_Chapters_1_to_5_Revised.pdf'


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Times-Roman', 9)
    canvas.drawCentredString(letter[0] / 2, 0.48 * inch, f'ToothWise Chapters IV and V | {doc.page}')
    canvas.restoreState()


def cell_text(cell):
    return '<br/>'.join(escape(p.text).replace('\n', '<br/>') for p in cell.paragraphs if p.text.strip()) or ' '


def main():
    source = SOURCE if SOURCE.exists() else ROOT / 'toothwise-updated-real-chapter-1-3-midtermpdfversion.pdf'
    if not source.exists():
        raise FileNotFoundError('Source Chapters I-III PDF not found')

    styles = getSampleStyleSheet()
    body = ParagraphStyle('Body', parent=styles['BodyText'], fontName='Times-Roman', fontSize=11,
                          leading=16, alignment=TA_JUSTIFY, firstLineIndent=0.25 * inch, spaceAfter=7)
    h1 = ParagraphStyle('H1', parent=styles['Heading1'], fontName='Times-Bold', fontSize=13,
                        leading=16, alignment=TA_CENTER, spaceBefore=13, spaceAfter=7)
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontName='Times-Bold', fontSize=11.5,
                        leading=14, spaceBefore=12, spaceAfter=5)
    cover = ParagraphStyle('Cover', parent=styles['Title'], fontName='Times-Bold', fontSize=16,
                           leading=20, alignment=TA_CENTER, spaceAfter=12)
    center = ParagraphStyle('Center', parent=body, alignment=TA_CENTER, firstLineIndent=0, spaceAfter=9)
    cell = ParagraphStyle('Cell', parent=body, fontSize=8.5, leading=10, firstLineIndent=0, spaceAfter=0)
    headcell = ParagraphStyle('HeadCell', parent=cell, textColor=colors.white, fontName='Times-Bold', alignment=TA_CENTER)

    document = Document(DOCX)
    story = []
    first = True
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            para = next((p for p in document.paragraphs if p._p is element), None)
            if para is None or not para.text.strip():
                continue
            text = escape(para.text).replace('\n', '<br/>')
            style_name = para.style.name
            if first:
                story += [Spacer(1, 1.2 * inch), Paragraph(text, cover), Spacer(1, .25 * inch)]
                first = False
            elif para.text == 'CHAPTER V':
                story += [PageBreak(), Paragraph(text, h1)]
            elif style_name == 'Heading 1':
                story.append(Paragraph(text, h1))
            elif style_name == 'Heading 2':
                story.append(Paragraph(text, h2))
            elif style_name.startswith('List Bullet'):
                story.append(Paragraph('&bull; ' + text, ParagraphStyle('Bullet', parent=body, firstLineIndent=0, leftIndent=.25*inch)))
            else:
                story.append(Paragraph(text, body))
        elif tag == 'tbl':
            table = next((t for t in document.tables if t._tbl is element), None)
            if table is None:
                continue
            data=[]
            for ri,row in enumerate(table.rows):
                style = headcell if ri == 0 else cell
                data.append([Paragraph(cell_text(c), style) for c in row.cells])
            widths = [6.55 * inch / len(data[0])] * len(data[0])
            rt = Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
            rt.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
                ('GRID', (0,0), (-1,-1), .35, colors.HexColor('#D9D9D9')),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('LEFTPADDING', (0,0), (-1,-1), 5), ('RIGHTPADDING', (0,0), (-1,-1), 5),
                ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F3F7FA')]),
            ]))
            story += [Spacer(1, 3), rt, Spacer(1, 5)]

    pdf = SimpleDocTemplate(str(CHAPTER_PDF), pagesize=letter, leftMargin=.9*inch, rightMargin=.9*inch,
                            topMargin=.85*inch, bottomMargin=.75*inch, title='ToothWise Chapters IV and V')
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)

    writer = PdfWriter()
    for path in (source, CHAPTER_PDF):
        reader = PdfReader(str(path))
        for page in reader.pages:
            writer.add_page(page)
    writer.add_metadata({'/Title': 'ToothWise Complete Chapters I to V Revised'})
    with open(OUTPUT, 'wb') as f:
        writer.write(f)
    print(OUTPUT)


if __name__ == '__main__':
    main()
